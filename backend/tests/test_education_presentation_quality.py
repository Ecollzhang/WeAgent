from io import BytesIO

import pytest
from pptx import Presentation

from services.edu.content_exporters import export_pptx
from services.edu.presentation_quality import (
    PRESENTATION_THEMES,
    inspect_pptx_bytes,
    inspect_slide_document,
)


def _document(style="clear_classroom"):
    return {
        "title": "校园叙事阅读与写作",
        "theme": {"style": style},
        "slides": [
            {
                "id": "slide-objectives",
                "title": "从细节读懂人物情绪",
                "layout": "content",
                "blocks": [
                    {
                        "type": "list",
                        "content": [
                            "定位时间、动作与心理描写",
                            "用文本证据解释人物情绪变化",
                            "迁移到一段校园生活续写",
                        ],
                    }
                ],
                "speaker_notes": "先让学生圈画证据，再两人互证。",
            },
            {
                "id": "slide-practice",
                "title": "证据链练习",
                "layout": "two_column",
                "blocks": [
                    {"type": "heading", "content": "Read"},
                    {
                        "type": "text",
                        "content": "Find the action that reveals the turning point.",
                    },
                    {"type": "heading", "content": "Write"},
                    {
                        "type": "text",
                        "content": "Continue the story in 80 words with one vivid detail.",
                    },
                ],
                "speaker_notes": "展示一份匿名样例进行同伴评价。",
            },
        ],
    }


def _font_sizes(shape):
    values = []
    if not getattr(shape, "has_text_frame", False):
        return values
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                values.append(run.font.size.pt)
    return values


def test_four_named_presentation_themes_are_available():
    assert {
        "clear_classroom": "清朗课堂",
        "paper_annotation": "纸张批注",
        "storybook": "童趣绘本",
        "dark_focus": "深色聚焦",
    } == {
        key: PRESENTATION_THEMES[key]["label"]
        for key in PRESENTATION_THEMES
    }


@pytest.mark.parametrize("style", list((
    "clear_classroom",
    "paper_annotation",
    "storybook",
    "dark_focus",
)))
def test_pptx_export_is_editable_readable_and_theme_specific(style):
    payload = export_pptx(_document(style), "fallback")
    deck = Presentation(BytesIO(payload))

    assert len(deck.slides) == 3
    assert deck.core_properties.subject == style
    assert any(
        run.font.name == "Noto Sans SC"
        for shape in deck.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    )

    cover_sizes = [
        size
        for shape in deck.slides[0].shapes
        for size in _font_sizes(shape)
        if shape.has_text_frame and "校园叙事" in shape.text
    ]
    assert cover_sizes and min(cover_sizes) >= 50

    content_title_sizes = [
        size
        for shape in deck.slides[1].shapes
        for size in _font_sizes(shape)
        if shape.has_text_frame and "从细节读懂" in shape.text
    ]
    assert content_title_sizes and min(content_title_sizes) >= 35

    report = inspect_pptx_bytes(payload)
    assert report["status"] == "passed", report
    assert report["slide_count"] == 3
    assert report["editable_text_shape_count"] >= 6
    assert not report["findings"]


def test_slide_document_quality_rejects_unreadable_density():
    dense = _document()
    dense["slides"][0]["blocks"] = [
        {"type": "text", "content": "过密内容" * 150},
        {"type": "list", "content": [f"第 {index} 项" for index in range(14)]},
    ]

    report = inspect_slide_document(dense)

    assert report["status"] == "failed"
    assert report["slide_count"] == 2
    assert any(
        finding["code"] == "slide_text_too_dense"
        for finding in report["slides"][0]["findings"]
    )
    assert any(
        finding["code"] == "too_many_list_items"
        for finding in report["slides"][0]["findings"]
    )


def test_slide_document_quality_accepts_a_teachable_deck():
    report = inspect_slide_document(_document("paper_annotation"))

    assert report["status"] == "passed", report
    assert report["theme_style"] == "paper_annotation"
    assert report["slide_count"] == 2
    assert all(item["status"] == "passed" for item in report["slides"])
