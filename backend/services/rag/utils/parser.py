"""文档解析 — 从各格式文件提取纯文本."""
import os
import io
import sys

# 可选依赖，安装时才生效
_PDF_OK = False
_DOCX_OK = False

try:
    import pdfplumber
    _PDF_OK = True
except ImportError:
    pass

try:
    import docx
    _DOCX_OK = True
except ImportError:
    pass

EXT_MAP = {
    '.pdf': 'pdf',
    '.txt': 'txt',
    '.md': 'md',
    '.markdown': 'md',
    '.docx': 'docx',
    '.csv': 'csv',
    '.json': 'json',
    '.py': 'code',
    '.js': 'code',
    '.ts': 'code',
    '.html': 'html',
    '.htm': 'html',
    '.xml': 'xml',
    '.yaml': 'yaml',
    '.yml': 'yaml',
}


def detect_file_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    return EXT_MAP.get(ext, 'unknown')


def extract_text(file_bytes: bytes, filename: str) -> tuple:
    """从文件二进制内容提取纯文本.

    Returns:
        (text: str, page_count: int, metadata: dict)
    """
    file_type = detect_file_type(filename)

    if file_type == 'pdf' and _PDF_OK:
        return _extract_pdf(file_bytes)
    elif file_type == 'txt' or file_type == 'md' or file_type == 'code':
        return _extract_text(file_bytes)
    elif file_type == 'docx' and _DOCX_OK:
        return _extract_docx(file_bytes)
    elif file_type in ('csv', 'json', 'yaml', 'xml', 'html'):
        text = file_bytes.decode('utf-8', errors='replace')
        return text, 1, {'file_type': file_type}
    else:
        # fallback: try as plain text
        try:
            text = file_bytes.decode('utf-8')
            return text, 1, {'file_type': 'text', 'warning': f'unrecognized extension, treated as text'}
        except UnicodeDecodeError:
            return '', 0, {'file_type': 'binary', 'error': f'unsupported file type: {file_type}'}


def _extract_pdf(file_bytes: bytes) -> tuple:
    text_parts = []
    page_count = 0
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
            page_count += 1
    return '\n\n'.join(text_parts), page_count, {'file_type': 'pdf'}


def _extract_text(file_bytes: bytes) -> tuple:
    text = file_bytes.decode('utf-8', errors='replace')
    return text, 1, {'file_type': 'text'}


def _extract_docx(file_bytes: bytes) -> tuple:
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = []
    for para in document.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)
    # also extract tables
    table_count = 0
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            paragraphs.append(' | '.join(cells))
        table_count += 1
    return '\n\n'.join(paragraphs), len(document.paragraphs), {
        'file_type': 'docx',
        'table_count': table_count,
    }
