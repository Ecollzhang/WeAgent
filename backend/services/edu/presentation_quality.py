"""Deterministic visual-quality checks for canonical Education presentations."""

from io import BytesIO

from .presentation_themes import (
    DEFAULT_PRESENTATION_THEME,
    PRESENTATION_THEMES,
    presentation_theme,
)


def _plain_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float, bool)):
        return str(value)
    if isinstance(value, list):
        return "\n".join(filter(None, (_plain_text(item) for item in value)))
    if isinstance(value, dict):
        return "\n".join(filter(None, (_plain_text(item) for item in value.values())))
    return str(value)


def _finding(code, message, severity="error"):
    return {"code": code, "message": message, "severity": severity}


def inspect_slide_document(source):
    """Inspect semantic density before any renderer receives the document."""
    source = source if isinstance(source, dict) else {}
    raw_theme = source.get("theme") if isinstance(source.get("theme"), dict) else {}
    requested_style = str(
        raw_theme.get("style")
        or raw_theme.get("preset")
        or raw_theme.get("id")
        or DEFAULT_PRESENTATION_THEME
    ).strip()
    style, _ = presentation_theme(source)
    slides = source.get("slides") if isinstance(source.get("slides"), list) else []
    deck_findings = []
    slide_reports = []

    if requested_style not in PRESENTATION_THEMES:
        deck_findings.append(
            _finding(
                "unknown_theme",
                f"未知课件风格 {requested_style}，导出时将使用清朗课堂。",
                "warning",
            )
        )
    if not str(source.get("title") or "").strip():
        deck_findings.append(_finding("missing_deck_title", "课件缺少标题。"))
    if not slides:
        deck_findings.append(_finding("missing_slides", "课件至少需要一页幻灯片。"))

    for index, slide in enumerate(slides):
        slide = slide if isinstance(slide, dict) else {}
        findings = []
        title = str(slide.get("title") or "").strip()
        blocks = slide.get("blocks") if isinstance(slide.get("blocks"), list) else []
        block_texts = [_plain_text(block.get("content")) for block in blocks if isinstance(block, dict)]
        text = "\n".join(filter(None, block_texts))
        list_item_count = sum(
            len(block.get("content"))
            for block in blocks
            if isinstance(block, dict) and isinstance(block.get("content"), list)
        )

        if not title:
            findings.append(_finding("missing_slide_title", "页面缺少标题。"))
        if len(title) > 72:
            findings.append(
                _finding(
                    "slide_title_too_long",
                    "页面标题超过 72 个字符，建议压缩为单一信息。",
                    "warning",
                )
            )
        if not blocks or not text:
            findings.append(_finding("empty_slide", "页面没有可见正文。"))
        if len(text) > 560:
            findings.append(
                _finding(
                    "slide_text_too_dense",
                    "页面正文超过 560 个字符，应拆分为多页。",
                )
            )
        elif len(text) > 400:
            findings.append(
                _finding(
                    "slide_text_dense",
                    "页面正文偏多，建议精简或拆页。",
                    "warning",
                )
            )
        if len(blocks) > 8:
            findings.append(
                _finding(
                    "too_many_blocks",
                    "页面包含超过 8 个内容块，应重新编排。",
                )
            )
        if list_item_count > 10:
            findings.append(
                _finding(
                    "too_many_list_items",
                    "单页列表超过 10 项，应拆分为多页。",
                )
            )
        for block in blocks:
            if not isinstance(block, dict) or block.get("type") != "image":
                continue
            content = block.get("content")
            content = content if isinstance(content, dict) else {}
            if not str(content.get("alt") or "").strip():
                findings.append(
                    _finding(
                        "image_missing_alt",
                        "图片内容块缺少替代文本。",
                        "warning",
                    )
                )

        status = (
            "failed"
            if any(item["severity"] == "error" for item in findings)
            else ("warning" if findings else "passed")
        )
        slide_reports.append(
            {
                "id": str(slide.get("id") or f"slide-{index + 1}"),
                "number": index + 1,
                "title": title or f"第 {index + 1} 页",
                "status": status,
                "metrics": {
                    "character_count": len(text),
                    "block_count": len(blocks),
                    "list_item_count": list_item_count,
                },
                "findings": findings,
            }
        )

    has_errors = any(item["severity"] == "error" for item in deck_findings) or any(
        slide["status"] == "failed" for slide in slide_reports
    )
    has_warnings = bool(deck_findings) or any(
        slide["status"] == "warning" for slide in slide_reports
    )
    return {
        "status": "failed" if has_errors else ("warning" if has_warnings else "passed"),
        "theme_style": style,
        "theme_label": PRESENTATION_THEMES[style]["label"],
        "slide_count": len(slides),
        "findings": deck_findings,
        "slides": slide_reports,
    }


def inspect_pptx_bytes(payload):
    """Check the actual PPTX package for editability and slide-bound violations."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        return {
            "status": "failed",
            "slide_count": 0,
            "editable_text_shape_count": 0,
            "findings": [
                _finding("pptx_adapter_unavailable", str(exc)),
            ],
            "slides": [],
        }

    try:
        deck = Presentation(BytesIO(payload))
    except Exception as exc:
        return {
            "status": "failed",
            "slide_count": 0,
            "editable_text_shape_count": 0,
            "findings": [_finding("invalid_pptx", f"PPTX 无法打开：{exc}")],
            "slides": [],
        }

    width, height = deck.slide_width, deck.slide_height
    editable_count = 0
    slide_reports = []
    all_findings = []
    tolerance = 5000
    for index, slide in enumerate(deck.slides, 1):
        findings = []
        visible_text = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                editable_count += 1
                if shape.text.strip():
                    visible_text.append(shape.text.strip())
            if (
                shape.left < -tolerance
                or shape.top < -tolerance
                or shape.left + shape.width > width + tolerance
                or shape.top + shape.height > height + tolerance
            ):
                findings.append(
                    _finding(
                        "shape_out_of_bounds",
                        f"形状 {shape.name} 超出页面边界。",
                    )
                )
        if not visible_text:
            findings.append(_finding("rendered_slide_empty", "导出的页面没有可编辑文字。"))
        status = "failed" if findings else "passed"
        slide_reports.append(
            {
                "number": index,
                "status": status,
                "editable_text_shape_count": sum(
                    1 for shape in slide.shapes if getattr(shape, "has_text_frame", False)
                ),
                "findings": findings,
            }
        )
        all_findings.extend(
            {"slide_number": index, **finding} for finding in findings
        )

    return {
        "status": "failed" if all_findings else "passed",
        "slide_count": len(deck.slides),
        "editable_text_shape_count": editable_count,
        "findings": all_findings,
        "slides": slide_reports,
    }
